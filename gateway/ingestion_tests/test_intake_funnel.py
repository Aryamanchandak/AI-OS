from __future__ import annotations

from pathlib import Path
import shutil

from docx import Document
import pymupdf
import pytest

from gateway.ingestion.intake.base import IntakeResult
from gateway.ingestion.intake.extractors import (
    LegacyWordExtractor,
    MarkdownPassthroughExtractor,
    PdfExtractor,
    WordExtractor,
)
from gateway.ingestion.intake.extractors.legacy_word import DocConversionError
from gateway.ingestion.intake.funnel import (
    DownstreamUnavailableError,
    UnroutableFileError,
    route_file,
    run_funnel,
    scan_drop_folder,
)


def _write_native_pdf(path: Path, text: str) -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(path)
    document.close()


def _write_scanned_pdf(path: Path) -> None:
    document = pymupdf.open()
    page = document.new_page()
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 100), False)
    pixmap.clear_with(255)
    page.insert_image(page.rect, stream=pixmap.tobytes("png"))
    document.save(path)
    document.close()


def test_scan_is_sorted_non_recursive_and_route_is_explicit(tmp_path: Path) -> None:
    (tmp_path / "b.docx").touch()
    (tmp_path / "A.PDF").touch()
    (tmp_path / "c.md").touch()
    (tmp_path / "book.txt").touch()
    (tmp_path / ".book.json").touch()
    (tmp_path / "processed").mkdir()
    (tmp_path / "processed" / "ignored.md").touch()

    paths = scan_drop_folder(tmp_path)
    assert [path.name for path in paths] == ["A.PDF", "b.docx", "c.md"]

    extractors = [PdfExtractor(), WordExtractor(), MarkdownPassthroughExtractor()]
    assert isinstance(route_file(paths[0], extractors), PdfExtractor)
    assert isinstance(route_file(paths[1], extractors), WordExtractor)
    assert isinstance(route_file(paths[2], extractors), MarkdownPassthroughExtractor)

    unknown = tmp_path / "unknown.txt"
    unknown.touch()
    with pytest.raises(UnroutableFileError, match="unknown.txt"):
        route_file(unknown, extractors)


def test_native_pdf_uses_text_layer_without_ocr(tmp_path: Path) -> None:
    path = tmp_path / "native.pdf"
    expected = "A digital PDF has exact embedded text for extraction."
    _write_native_pdf(path, expected)

    def unexpected_ocr(_pixmap: pymupdf.Pixmap) -> list[dict]:
        raise AssertionError("native PDF must not use OCR")

    result = PdfExtractor(ocr_lines=unexpected_ocr).extract(path)
    assert expected in result.markdown_text
    assert result.needs_approval is False
    assert result.frontmatter["extraction_method"] == "pdf_text_layer"


def test_scanned_pdf_delegates_to_ocr_and_requires_approval(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    _write_scanned_pdf(path)
    calls = []

    def fake_ocr(pixmap: pymupdf.Pixmap) -> list[dict]:
        calls.append(pixmap)
        return [
            {"text": "Recognized scanned page text with enough content.", "confidence": 0.9}
        ]

    result = PdfExtractor(ocr_lines=fake_ocr).extract(path)
    assert len(calls) == 1
    assert "Recognized scanned page" in result.markdown_text
    assert result.needs_approval is True
    assert "required OCR: 1" in (result.approval_reason or "")
    assert result.frontmatter["ocr_pages"] == [1]


def test_word_extracts_paragraphs_and_corruption_fails_loud(tmp_path: Path) -> None:
    path = tmp_path / "document.docx"
    document = Document()
    document.add_paragraph("First paragraph with enough text.")
    document.add_paragraph("Second paragraph is preserved too.")
    document.save(path)

    result = WordExtractor().extract(path)
    assert result.markdown_text == (
        "First paragraph with enough text.\n\nSecond paragraph is preserved too."
    )
    assert result.needs_approval is False

    corrupt = tmp_path / "corrupt.docx"
    corrupt.write_bytes(b"not an OOXML archive")
    with pytest.raises(Exception):
        WordExtractor().extract(corrupt)


def test_word_preserves_table_content_and_reading_order(tmp_path: Path) -> None:
    path = tmp_path / "with_table.docx"
    document = Document()
    document.add_paragraph("Intro paragraph before the table.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "1"
    document.add_paragraph("Outro paragraph after the table.")
    document.save(path)

    result = WordExtractor().extract(path)

    assert "Intro paragraph before the table." in result.markdown_text
    assert "Outro paragraph after the table." in result.markdown_text
    assert "| Name | Value |" in result.markdown_text
    assert "| alpha | 1 |" in result.markdown_text
    intro_index = result.markdown_text.index("Intro paragraph")
    table_index = result.markdown_text.index("| Name | Value |")
    outro_index = result.markdown_text.index("Outro paragraph")
    assert intro_index < table_index < outro_index
    assert result.needs_approval is False


def test_word_maps_heading_and_list_styles_to_markdown(tmp_path: Path) -> None:
    path = tmp_path / "styled.docx"
    document = Document()
    document.add_paragraph("Section One", style="Heading 1")
    document.add_paragraph("First bullet point here.", style="List Bullet")
    document.add_paragraph("Second bullet point here.", style="List Bullet")
    document.save(path)

    result = WordExtractor().extract(path)

    assert "# Section One" in result.markdown_text
    assert "- First bullet point here." in result.markdown_text
    assert "- Second bullet point here." in result.markdown_text


def test_word_uses_core_property_title_when_present(tmp_path: Path) -> None:
    path = tmp_path / "titled.docx"
    document = Document()
    document.core_properties.title = "Real Document Title"
    document.add_paragraph("Body text with enough length to pass sanity checks.")
    document.save(path)

    result = WordExtractor().extract(path)

    assert result.title == "Real Document Title"


def test_word_flags_embedded_images_for_approval(tmp_path: Path) -> None:
    import io

    from PIL import Image

    path = tmp_path / "with_image.docx"
    document = Document()
    document.add_paragraph("Body text with enough length to pass sanity checks.")
    image_bytes = io.BytesIO()
    Image.new("RGB", (2, 2), color="red").save(image_bytes, format="PNG")
    image_bytes.seek(0)
    document.add_picture(image_bytes)
    document.save(path)

    result = WordExtractor().extract(path)

    assert result.needs_approval is True
    assert "embedded image" in (result.approval_reason or "")
    assert result.frontmatter["embedded_image_count"] == 1


def test_legacy_doc_converts_via_textutil_and_reuses_docx_parsing(
    tmp_path: Path,
) -> None:
    path = tmp_path / "legacy.doc"

    def fake_convert(source: Path, destination: Path) -> None:
        assert source == path
        document = Document()
        document.add_paragraph("Converted from a legacy binary .doc file.")
        document.save(destination)

    result = LegacyWordExtractor(
        convert_to_docx=fake_convert, detect_tables=lambda _path: False
    ).extract(path)

    assert result.markdown_text == "Converted from a legacy binary .doc file."
    assert result.needs_approval is False
    assert result.frontmatter["source_filename"] == "legacy.doc"
    assert result.frontmatter["extraction_method"] == "doc_textutil_conversion"


def test_legacy_doc_conversion_failure_fails_loud(tmp_path: Path) -> None:
    path = tmp_path / "legacy.doc"

    def failing_convert(source: Path, destination: Path) -> None:
        raise DocConversionError("textutil exploded")

    with pytest.raises(DocConversionError, match="textutil exploded"):
        LegacyWordExtractor(
            convert_to_docx=failing_convert, detect_tables=lambda _path: False
        ).extract(path)


def test_legacy_doc_with_detected_table_requires_approval(tmp_path: Path) -> None:
    path = tmp_path / "legacy.doc"

    def fake_convert(source: Path, destination: Path) -> None:
        document = Document()
        document.add_paragraph("Name")
        document.add_paragraph("alpha")
        document.save(destination)

    result = LegacyWordExtractor(
        convert_to_docx=fake_convert, detect_tables=lambda _path: True
    ).extract(path)

    assert result.needs_approval is True
    assert "table" in (result.approval_reason or "")
    assert "flatten" in (result.approval_reason or "")


def test_legacy_doc_inconclusive_table_detection_requires_approval(
    tmp_path: Path,
) -> None:
    path = tmp_path / "legacy.doc"

    def fake_convert(source: Path, destination: Path) -> None:
        document = Document()
        document.add_paragraph("Body text with enough length to pass sanity checks.")
        document.save(destination)

    result = LegacyWordExtractor(
        convert_to_docx=fake_convert, detect_tables=lambda _path: None
    ).extract(path)

    assert result.needs_approval is True
    assert "could not verify" in (result.approval_reason or "")


def test_legacy_doc_handles_suffix_and_routes_ahead_of_docx(tmp_path: Path) -> None:
    doc_path = tmp_path / "old.doc"
    docx_path = tmp_path / "new.docx"
    doc_path.touch()
    docx_path.touch()

    legacy = LegacyWordExtractor()
    word = WordExtractor()
    assert legacy.handles(doc_path) is True
    assert legacy.handles(docx_path) is False
    assert word.handles(doc_path) is False
    assert word.handles(docx_path) is True

    extractors = [word, legacy]
    assert isinstance(route_file(doc_path, extractors), LegacyWordExtractor)
    assert isinstance(route_file(docx_path, extractors), WordExtractor)


@pytest.mark.skipif(
    shutil.which("textutil") is None, reason="requires macOS's textutil"
)
def test_legacy_doc_real_textutil_conversion_end_to_end(tmp_path: Path) -> None:
    import subprocess

    seed = tmp_path / "seed.txt"
    seed.write_text("Body text produced through a real textutil round trip.")
    doc_path = tmp_path / "real.doc"
    subprocess.run(
        ["textutil", "-convert", "doc", "-output", str(doc_path), str(seed)],
        check=True,
        capture_output=True,
    )

    result = LegacyWordExtractor().extract(doc_path)

    assert "Body text produced through a real textutil round trip." in (
        result.markdown_text
    )
    assert result.needs_approval is False
    assert result.frontmatter["source_filename"] == "real.doc"
    assert result.frontmatter["extraction_method"] == "doc_textutil_conversion"


@pytest.mark.skipif(
    shutil.which("textutil") is None, reason="requires macOS's textutil"
)
def test_legacy_doc_real_table_flattening_is_detected_and_flagged(
    tmp_path: Path,
) -> None:
    import subprocess

    seed_docx = tmp_path / "seed_with_table.docx"
    seed_document = Document()
    seed_document.add_paragraph("Intro before table")
    table = seed_document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "1"
    seed_document.save(seed_docx)

    doc_path = tmp_path / "with_table.doc"
    subprocess.run(
        ["textutil", "-convert", "doc", "-output", str(doc_path), str(seed_docx)],
        check=True,
        capture_output=True,
    )

    result = LegacyWordExtractor().extract(doc_path)

    # textutil's .doc -> .docx conversion really does flatten the table:
    # the cell text survives (nothing silently vanishes) but arrives as
    # plain paragraphs, not a markdown table -- confirmed by inspecting the
    # converted document directly, not assumed.
    assert "Name" in result.markdown_text
    assert "alpha" in result.markdown_text
    assert "|" not in result.markdown_text
    assert result.needs_approval is True
    assert "flatten" in (result.approval_reason or "")


def test_markdown_is_unchanged_and_invalid_utf8_fails(tmp_path: Path) -> None:
    path = tmp_path / "note.md"
    content = "# Exact note\n\nThis content passes through byte-for-byte.\n"
    path.write_text(content, encoding="utf-8")
    assert MarkdownPassthroughExtractor().extract(path).markdown_text == content

    invalid = tmp_path / "invalid.md"
    invalid.write_bytes(b"\xff\xfe")
    with pytest.raises(UnicodeDecodeError):
        MarkdownPassthroughExtractor().extract(invalid)


class FixtureExtractor:
    def handles(self, path: Path) -> bool:
        return path.suffix == ".fixture"

    def extract(self, path: Path) -> IntakeResult:
        requires_review = path.stem.startswith("review")
        return IntakeResult(
            markdown_text="Extracted fixture content long enough to pass checks.",
            title=path.stem,
            needs_approval=requires_review,
            approval_reason="fixture requires review" if requires_review else None,
        )


def test_run_funnel_publishes_or_parks_without_crossing_the_gate(tmp_path: Path) -> None:
    drop = tmp_path / "drop"
    vault = tmp_path / "vault"
    drop.mkdir()
    vault.mkdir()
    approved = drop / "approved.fixture"
    review = drop / "review-me.fixture"
    approved.write_text("approved", encoding="utf-8")
    review.write_text("review", encoding="utf-8")
    writes: list[str] = []
    staged: list[tuple[Path, str]] = []

    def writer(vault_path: Path, result: IntakeResult) -> Path:
        writes.append(result.title)
        destination = vault_path / f"{result.title}.md"
        destination.write_text(result.markdown_text, encoding="utf-8")
        return destination

    def stager(source: Path, result: IntakeResult) -> int:
        staged.append((source, result.title))
        return 17

    report = run_funnel(
        drop,
        vault,
        extractors=[FixtureExtractor()],
        write_note_to_vault=writer,
        park_for_review=stager,
    )

    assert report.published_count == 1
    assert report.awaiting_approval_count == 1
    assert writes == ["approved"]
    assert staged == [(review, "review-me")]
    assert (drop / "processed" / "approved.fixture").exists()
    assert review.exists()
    assert not (vault / "review-me.md").exists()


def test_run_funnel_never_moves_when_required_handoff_is_missing(tmp_path: Path) -> None:
    drop = tmp_path / "drop"
    vault = tmp_path / "vault"
    drop.mkdir()
    vault.mkdir()
    source = drop / "approved.fixture"
    source.touch()

    with pytest.raises(DownstreamUnavailableError, match="vault writer"):
        run_funnel(drop, vault, extractors=[FixtureExtractor()])
    assert source.exists()


def test_run_funnel_does_not_move_a_file_when_publish_fails(tmp_path: Path) -> None:
    drop = tmp_path / "drop"
    vault = tmp_path / "vault"
    drop.mkdir()
    vault.mkdir()
    source = drop / "approved.fixture"
    source.touch()

    def failed_writer(_vault: Path, _result: IntakeResult) -> None:
        raise RuntimeError("vault unavailable")

    with pytest.raises(RuntimeError, match="vault unavailable"):
        run_funnel(
            drop,
            vault,
            extractors=[FixtureExtractor()],
            write_note_to_vault=failed_writer,
        )
    assert source.exists()
    assert not (drop / "processed").exists()


def test_run_funnel_supports_extractor_specific_post_publish_destination(
    tmp_path: Path,
) -> None:
    drop = tmp_path / "drop"
    vault = tmp_path / "vault"
    archive = tmp_path / "archive"
    drop.mkdir()
    vault.mkdir()
    source = drop / "approved.fixture"
    source.touch()

    def writer(_vault: Path, _result: IntakeResult) -> Path:
        return vault / "approved.md"

    def archive_destination(path: Path, result: IntakeResult) -> Path:
        assert path == source
        assert result.title == "approved"
        return archive / path.name

    run_funnel(
        drop,
        vault,
        extractors=[FixtureExtractor()],
        write_note_to_vault=writer,
        post_publish_destination={FixtureExtractor: archive_destination},
    )

    assert (archive / source.name).exists()
    assert not (drop / "processed").exists()

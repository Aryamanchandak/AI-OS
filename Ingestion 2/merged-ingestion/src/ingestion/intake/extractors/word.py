from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from ingestion.intake.base import IntakeResult, text_sanity_issue

_HEADING_STYLE_LEVELS = {f"Heading {level}": level for level in range(1, 7)}
_BULLET_LIST_STYLES = {"List Bullet", "List Bullet 2", "List Bullet 3"}
_NUMBER_LIST_STYLES = {"List Number", "List Number 2", "List Number 3"}


def _iter_block_items(document: DocxDocument) -> list[Paragraph | Table]:
    """Yield body paragraphs and tables in document order.

    python-docx does not expose this directly: ``document.paragraphs`` and
    ``document.tables`` are separate flat lists that lose interleaving and,
    for tables, drop the cell text entirely if you only read ``paragraphs``.
    This walks the underlying body XML instead, per the standard python-docx
    recipe for ordered block iteration.
    """

    items: list[Paragraph | Table] = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            items.append(Paragraph(child, document))
        elif isinstance(child, CT_Tbl):
            items.append(Table(child, document))
    return items


def _paragraph_to_markdown(paragraph: Paragraph) -> str | None:
    text = paragraph.text.strip()
    if not text:
        return None
    style_name = paragraph.style.name if paragraph.style else ""
    heading_level = _HEADING_STYLE_LEVELS.get(style_name)
    if heading_level:
        return f"{'#' * heading_level} {text}"
    if style_name in _BULLET_LIST_STYLES:
        return f"- {text}"
    if style_name in _NUMBER_LIST_STYLES:
        return f"1. {text}"
    return text


def _table_to_markdown(table: Table) -> str | None:
    def cell_text(raw: str) -> str:
        return raw.replace("\n", " ").replace("|", "\\|").strip()

    rows = [[cell_text(cell.text) for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return None

    header, *body_rows = rows
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body_rows)
    return "\n".join(lines)


def extract_docx_document(
    document: DocxDocument,
    *,
    source_path: Path,
    extraction_method: str,
    extra_reasons: Sequence[str] = (),
) -> IntakeResult:
    """Build an IntakeResult from an already-open OOXML document.

    Shared by ``WordExtractor`` (native .docx) and the legacy .doc extractor,
    which converts to .docx first and then reuses this exact logic so both
    paths get identical table/heading/image handling. ``extra_reasons`` lets
    a caller fold in approval reasons discovered before this document was
    even opened (e.g. the .doc conversion step detecting lost fidelity).
    """

    blocks: list[str] = []
    for item in _iter_block_items(document):
        rendered = (
            _paragraph_to_markdown(item)
            if isinstance(item, Paragraph)
            else _table_to_markdown(item)
        )
        if rendered:
            blocks.append(rendered)
    markdown_text = "\n\n".join(blocks)

    title = document.core_properties.title or source_path.stem

    reasons: list[str] = list(extra_reasons)
    sanity_issue = text_sanity_issue(markdown_text)
    if sanity_issue:
        reasons.append(sanity_issue)
    embedded_image_count = len(document.inline_shapes)
    if embedded_image_count:
        reasons.append(
            f"document contains {embedded_image_count} embedded image(s) "
            "that extraction does not capture"
        )

    return IntakeResult(
        markdown_text=markdown_text,
        title=title,
        needs_approval=bool(reasons),
        approval_reason="; ".join(reasons) or None,
        frontmatter={
            "source_filename": source_path.name,
            "extraction_method": extraction_method,
            "embedded_image_count": embedded_image_count,
        },
    )


class WordExtractor:
    """Extract paragraph and table text from an Office Open XML document."""

    def handles(self, path: Path) -> bool:
        return path.is_file() and path.suffix.lower() == ".docx"

    def extract(self, path: Path) -> IntakeResult:
        document = Document(str(path))
        return extract_docx_document(
            document, source_path=path, extraction_method="docx_xml"
        )


def extract_word(path: Path) -> IntakeResult:
    return WordExtractor().extract(path)

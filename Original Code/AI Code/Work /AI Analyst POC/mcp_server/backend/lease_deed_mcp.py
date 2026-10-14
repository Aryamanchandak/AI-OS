"""
lease_deed.py
-------------
MCP tools: build_rent_roll, build_clause_report, process_lease_deed
"""

import json
import os
import re
from datetime import datetime, date
from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from google import genai
from google.genai import types

from mcp_server import mcp


# ─────────────────────────────────────────────────────────────────────────────
# RENT ROLL PROMPTS
# ─────────────────────────────────────────────────────────────────────────────

_RENT_ROLL_SYSTEM_PROMPT = """You are a commercial real estate lease abstraction expert.
Your job is to read a Leave & License Agreement (or any commercial lease deed)
and extract structured data for a rent roll.

IMPORTANT RULES:
1. A single lease deed may contain MORE THAN ONE contract/schedule
   (e.g. two different units with different carpet areas or leasable areas).
   Return ONE JSON object per contract — i.e. an ARRAY of objects.
2. If a field is not mentioned in the document, return null for that field.
3. Never invent or assume values. Only extract what is explicitly stated.
4. Dates must be returned as "YYYY-MM-DD" strings.
5. All monetary amounts and areas must be plain numbers (no ₹, Rs, commas).
6. Escalation % must be a decimal (5% → 0.05).
7. Efficiency % = carpet_area / leasable_area — calculate it yourself.
8. psf_pm values = monthly amount / leasable_area — calculate yourself.
9. For remaining_lease_period_yrs, return null — it will be computed live.
10. For remaining_lease_period_incl_renewal_yrs, return null — computed live.

CRITICAL JSON RULES:
- Return ONLY a valid JSON array. No markdown, no explanation, no code fences.
- Ensure all strings are properly JSON-escaped:
  * Use \\n for newlines (not literal line breaks inside strings)
  * Use \\" for quotes inside string values
  * No unescaped special characters
- Every opening { must have a matching closing }
- Every opening [ must have a matching closing ]
- No trailing commas before } or ]
- All string values must be wrapped in double quotes

Each object in the array must follow this exact schema:
{
  "sno": <integer, sequential starting 1>,
  "tenant_name": <string>,
  "block_and_floor": <string>,
  "status": <string, e.g. "Leased">,
  "carpet_area_sf": <number or null>,
  "leasable_area_sf": <number or null>,
  "efficiency_pct": <decimal 0-1 or null>,
  "condition": <string or null>,
  "base_rent_psf_pm": <number or null>,
  "current_rent_psf_pm": <number or null>,
  "monthly_rent": <number or null>,
  "fitout_rent_psf_pm": <number or null>,
  "monthly_fitout_income": <number or null>,
  "base_cam_psf_pm": <number or null>,
  "current_cam_psf_pm": <number or null>,
  "monthly_cam": <number or null>,
  "initial_lcd": <"YYYY-MM-DD" or null>,
  "current_lcd": <"YYYY-MM-DD" or null>,
  "rent_start_date": <"YYYY-MM-DD" or null>,
  "lock_in_period_yrs": <number or null>,
  "lock_in_end_date": <"YYYY-MM-DD" or null>,
  "lease_end_date": <"YYYY-MM-DD" or null>,
  "remaining_lease_period_yrs": null,
  "lease_end_date_incl_renewal": <"YYYY-MM-DD" or null>,
  "remaining_lease_period_incl_renewal_yrs": null,
  "renewal_terms": <string or null>,
  "renewal_option": <"Y" or "N" or null>,
  "security_deposit_inr": <number or null>,
  "security_deposit_received_inr": <number or null>,
  "paid_parking_slots": <number or null>,
  "parking_charges_for_paid_slots": <number or null>,
  "rent_escalation_pct": <decimal or null>,
  "period_years": <number or null>,
  "first_rent_escalation_date": <"YYYY-MM-DD" or null>,
  "first_escalation_lease_rent_psf_pm": <number or null>,
  "second_rent_escalation_date": <"YYYY-MM-DD" or null>,
  "second_escalation_lease_rent_psf_pm": <number or null>,
  "next_escalation_cam_psf_pm": <number or null>,
  "new_tenant_lcd": <"YYYY-MM-DD" or null>,
  "new_rent_psf_pm": <number or null>
}"""

_RENT_ROLL_USER_PROMPT = """Extract all contracts from this lease deed and return the JSON array.
Remember: one object per contract/unit. If there are two units with different areas, return two objects.
IMPORTANT: Return ONLY the JSON array with no markdown, no explanations, no code fences."""


# ─────────────────────────────────────────────────────────────────────────────
# CLAUSE ANALYSIS PROMPT
# ─────────────────────────────────────────────────────────────────────────────

_CLAUSE_ANALYSIS_PROMPT = """You are a senior legal analyst specialising in commercial real estate leases in India.

Your task: Carefully read the entire lease deed / Leave & License Agreement and identify ALL clauses that are critical and should be brought to people's attention.

Focus on clauses that, **if missed or overlooked**, could cause:

1. **MONETARY RISK** — Hidden penalties, ambiguous rent escalation, unilateral rent revision rights, unclear maintenance charges, TDS obligations, stamp duty liability, interest on delayed payments, forfeiture of security deposit, cost of restoration/fit-out removal, penalty for early exit, etc.

2. **DEAL RISK** — Unfavorable lock-in terms, auto-renewal without consent, one-sided termination rights, restrictive use clauses, subletting restrictions, change of ownership / assignment clauses, conditions precedent that may block the deal, right of first refusal, etc.

3. **LITIGATION RISK** — Vague or missing dispute resolution mechanism, unclear jurisdiction, absent indemnity clauses, missing force majeure, poorly defined default/breach/cure periods, lack of notice requirements, ambiguous handover/possession terms, unclear liability for structural defects, etc.

4. **OTHER IMPORTANT** — Any other clause that a prudent businessperson or legal reviewer should be aware of but does not fit neatly into the above three categories.

FOR EACH CLAUSE YOU IDENTIFY, return a JSON object with these exact fields:
{
  "category": "MONETARY_RISK" | "DEAL_RISK" | "LITIGATION_RISK" | "OTHER_IMPORTANT",
  "severity": "HIGH" | "MEDIUM" | "LOW",
  "clause_title": "<short descriptive title, e.g. 'Early Termination Penalty'>",
  "clause_text": "<verbatim excerpt from the document — copy the exact wording, keep it under 300 words>",
  "risk_summary": "<1-3 sentence plain-English explanation of WHY this clause is risky and what could go wrong>",
  "page_number": <integer page number where this clause appears, or null if uncertain>,
  "recommendation": "<brief actionable advice, e.g. 'Negotiate a mutual termination right with 90-day notice'>"
}

RULES:
- Return a JSON ARRAY of such objects. No markdown, no explanation, no code fences.
- Include at least 5 clauses if the document is substantial. There is no upper limit — be thorough.
- Severity guide:
  • HIGH = Could directly cause significant financial loss or legal dispute
  • MEDIUM = Creates moderate risk or unfavorable terms worth negotiating
  • LOW = Minor but worth noting for awareness
- Be specific. Quote the actual clause text from the document.
- If you cannot determine the page number, set it to null.
- Order clauses by severity (HIGH first, then MEDIUM, then LOW) within each category.

Return ONLY the JSON array."""


# ─────────────────────────────────────────────────────────────────────────────
# EXCEL BUILDER CONSTANTS
# ─────────────────────────────────────────────────────────────────────────────

_HEADERS = [
    "S.No.",                                                #  1
    "Tenant Name",                                          #  2
    "Block & Floor",                                        #  3
    "Status",                                               #  4
    "Carpet Area (sf)",                                     #  5
    "Leasable Area (sf)",                                   #  6
    "Efficiency %",                                         #  7
    "Condition",                                            #  8
    "Base Rent\nINR psf pm",                                #  9
    "Current Rent\nINR psf pm",                             # 10
    "Monthly Rent",                                         # 11
    "Fit Out Rent\nINR psf pm",                             # 12
    "Monthly Fitout Income",                                # 13
    "Base CAM\nINR psf pm",                                 # 14
    "Current CAM\nINR psf pm",                              # 15
    "Monthly CAM",                                          # 16
    "Initial LCD",                                          # 17
    "Current LCD",                                          # 18
    "Rent Start Date",                                      # 19
    "Lock In Period (Yrs)",                                 # 20
    "Lock In End Date",                                     # 21
    "Lease End Date",                                       # 22
    "Remaining Lease Period (Yrs)",                         # 23
    "Lease End Date\n(incl Renewal Option)",                # 24
    "Remaining Lease Period\n(incl Renewal Option) (Yrs)", # 25
    "Renewal Terms",                                        # 26
    "Renewal Option",                                       # 27
    "Security Deposit (INR)",                               # 28
    "Security Deposit Received (INR)",                      # 29
    "Paid Parking Slots",                                   # 30
    "Parking Charges for Paid Slots",                       # 31
    "Rent Escalation %",                                    # 32
    "Period (Years)",                                       # 33
    "First Rent Escalation Date",                           # 34
    "First Escalation Lease Rent\nINR psf pm",              # 35
    "",                                                     # 36 separator
    "Second Rent Escalation Date",                          # 37
    "Second Escalation Lease Rent\nINR psf pm",             # 38
    "Next Escalation CAM\nINR psf pm",                      # 39
    "New Tenant LCD",                                       # 40
    "New Rent\nINR psf pm",                                 # 41
]

_COL_MAP = {
    1:  "sno",
    2:  "tenant_name",
    3:  "block_and_floor",
    4:  "status",
    5:  "carpet_area_sf",
    6:  "leasable_area_sf",
    7:  "efficiency_pct",
    8:  "condition",
    9:  "base_rent_psf_pm",
    10: "current_rent_psf_pm",
    11: "monthly_rent",
    12: "fitout_rent_psf_pm",
    13: "monthly_fitout_income",
    14: "base_cam_psf_pm",
    15: "current_cam_psf_pm",
    16: "monthly_cam",
    17: "initial_lcd",
    18: "current_lcd",
    19: "rent_start_date",
    20: "lock_in_period_yrs",
    21: "lock_in_end_date",
    22: "lease_end_date",
    23: "remaining_lease_period_yrs",
    24: "lease_end_date_incl_renewal",
    25: "remaining_lease_period_incl_renewal_yrs",
    26: "renewal_terms",
    27: "renewal_option",
    28: "security_deposit_inr",
    29: "security_deposit_received_inr",
    30: "paid_parking_slots",
    31: "parking_charges_for_paid_slots",
    32: "rent_escalation_pct",
    33: "period_years",
    34: "first_rent_escalation_date",
    35: "first_escalation_lease_rent_psf_pm",
    36: None,   # separator
    37: "second_rent_escalation_date",
    38: "second_escalation_lease_rent_psf_pm",
    39: "next_escalation_cam_psf_pm",
    40: "new_tenant_lcd",
    41: "new_rent_psf_pm",
}

_DATE_COLS     = {17, 18, 19, 21, 22, 24, 34, 37, 40}
_PCT_COLS      = {7, 32}
_PSF_COLS      = {9, 10, 12, 14, 15, 35, 38, 39, 41}
_CURRENCY_COLS = {11, 13, 16, 28, 29}
_YRS_COLS      = {20, 23, 25, 33}
_INT_COLS      = {1, 5, 6, 30}

_COL_WIDTHS = [
    6, 22, 13, 10, 14, 14, 12, 13,
    16, 16, 14, 16, 18, 15, 16, 13,
    13, 13, 15, 14, 16, 15,
    24, 20, 30, 15, 14,
    22, 26,
    18, 26,
    16, 14,
    24, 26,
    3,
    24, 26,
    22, 16, 14,
]


# ─────────────────────────────────────────────────────────────────────────────
# WORD DOC BUILDER CONSTANTS
# ─────────────────────────────────────────────────────────────────────────────

_C_TITLE     = RGBColor(0x1E, 0x40, 0xAF)   # deep blue
_C_SECTION   = RGBColor(0x1D, 0x4E, 0xD8)   # blue
_C_BODY      = RGBColor(0x1F, 0x29, 0x37)   # near-black
_C_DESC      = RGBColor(0x4B, 0x55, 0x63)   # slate
_C_LABEL     = RGBColor(0x9C, 0xA3, 0xAF)   # light grey
_C_META      = RGBColor(0x6B, 0x72, 0x80)   # medium grey
_C_HIGH      = RGBColor(0xDC, 0x26, 0x26)   # red
_C_MEDIUM    = RGBColor(0xD9, 0x77, 0x06)   # amber
_C_LOW       = RGBColor(0x16, 0xA3, 0x4A)   # green
_C_CLAUSE_BG = RGBColor(0x37, 0x41, 0x51)   # dark grey for clause text
_C_REC       = RGBColor(0x05, 0x59, 0x88)   # teal for recommendations

_SEVERITY_COLORS = {"HIGH": _C_HIGH, "MEDIUM": _C_MEDIUM, "LOW": _C_LOW}
_SEVERITY_ICONS  = {"HIGH": "🔴", "MEDIUM": "🟡", "LOW": "🟢"}

_CATEGORY_TITLES = {
    "MONETARY_RISK":   "💰 Monetary Risk Clauses",
    "DEAL_RISK":       "🤝 Deal Risk Clauses",
    "LITIGATION_RISK": "⚖️ Litigation Risk Clauses",
    "OTHER_IMPORTANT": "📋 Other Important Clauses",
}

_CATEGORY_DESCRIPTIONS = {
    "MONETARY_RISK":   "Clauses that could cause direct or indirect financial loss if overlooked.",
    "DEAL_RISK":       "Clauses that could jeopardise the deal structure, timeline, or negotiating position.",
    "LITIGATION_RISK": "Clauses with ambiguous or missing terms that could lead to legal disputes.",
    "OTHER_IMPORTANT": "Additional clauses that a prudent reviewer should be aware of.",
}

_CATEGORY_ORDER = ["MONETARY_RISK", "DEAL_RISK", "LITIGATION_RISK", "OTHER_IMPORTANT"]


# ─────────────────────────────────────────────────────────────────────────────
# PRIVATE HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _parse_llm_json(raw: str) -> list[dict]:
    cleaned = re.sub(r"```(?:json)?|```", "", raw).strip()
    match = re.search(r"\[.*\]", cleaned, re.DOTALL)
    if not match:
        raise ValueError(f"No JSON array found in LLM response: {raw[:200]}")
    return json.loads(match.group())


def _extract_contracts(pdf_bytes: bytes) -> list[dict]:
    client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
    pdf_part = types.Part.from_bytes(data=pdf_bytes, mime_type="application/pdf")
    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=[
            types.Content(
                role="user",
                parts=[
                    types.Part.from_text(text=_RENT_ROLL_SYSTEM_PROMPT),
                    pdf_part,
                    types.Part.from_text(text=_RENT_ROLL_USER_PROMPT),
                ],
            )
        ],
    )
    try:
        return _parse_llm_json(response.text)
    except ValueError as e:
        raise ValueError(f"_extract_contracts: failed to parse Gemini response. {e}") from e


def _remaining_yrs(end_date_str: str | None) -> float | None:
    if not end_date_str:
        return None
    try:
        end = datetime.strptime(end_date_str, "%Y-%m-%d").date()
        return round((end - date.today()).days / 365.25, 2)
    except Exception:
        return None


def _enrich_contract(c: dict) -> dict:
    c["remaining_lease_period_yrs"] = _remaining_yrs(c.get("lease_end_date"))
    c["remaining_lease_period_incl_renewal_yrs"] = _remaining_yrs(
        c.get("lease_end_date_incl_renewal")
    )
    return c


def _parse_date(val) -> datetime | None:
    if not val:
        return None
    if isinstance(val, datetime):
        return val
    try:
        return datetime.strptime(str(val), "%Y-%m-%d")
    except Exception:
        return None


def _parse_clauses(raw: str) -> list[dict]:
    raw = re.sub(r"```(?:json)?|```", "", raw).strip()
    raw = re.sub(r',(\s*[}\]])', r'\1', raw)
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r'\[\s*\{.*?\}\s*\]', raw, re.DOTALL)
        if match:
            data = json.loads(match.group(0))
        else:
            raise ValueError(f"No JSON array found in response. First 300 chars: {raw[:300]}")
    if isinstance(data, dict):
        data = [data]
    clauses = []
    for item in data:
        if not isinstance(item, dict):
            continue
        item.setdefault("category", "OTHER_IMPORTANT")
        item.setdefault("severity", "MEDIUM")
        item.setdefault("clause_title", "Untitled Clause")
        item.setdefault("clause_text", "")
        item.setdefault("risk_summary", "")
        item.setdefault("page_number", None)
        item.setdefault("recommendation", "")
        clauses.append(item)
    return clauses


def _analyze_clauses(pdf_bytes: bytes) -> list[dict]:
    client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
    pdf_part = types.Part.from_bytes(data=pdf_bytes, mime_type="application/pdf")
    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=[
            types.Content(
                role="user",
                parts=[pdf_part, types.Part.from_text(text=_CLAUSE_ANALYSIS_PROMPT)],
            )
        ],
        config=types.GenerateContentConfig(
            thinking_config=types.ThinkingConfig(thinking_level="low"),
            response_mime_type="application/json",
        ),
    )
    try:
        return _parse_clauses(response.text)
    except ValueError as e:
        raise ValueError(f"_analyze_clauses: failed to parse Gemini response. {e}") from e


def _sanitize(text: str) -> str:
    if not text:
        return text
    text = text.replace("\x00", "")
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    return text


def _hr(doc, color="BFDBFE"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "4")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), color)
    pBdr.append(bot)
    pPr.append(pBdr)
    return p


def _build_docx(clauses: list[dict], document_name: str = "Lease Deed") -> bytes:
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    doc = Document()

    sec = doc.sections[0]
    sec.top_margin    = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin   = Inches(0.9)
    sec.right_margin  = Inches(0.9)

    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    # ── Title ─────────────────────────────────────────────────────────────────
    tp = doc.add_paragraph()
    r = tp.add_run("Critical Clause Analysis Report")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = _C_TITLE
    tp.paragraph_format.space_after = Pt(3)

    sp = doc.add_paragraph()
    sr = sp.add_run("AI-powered identification of high-risk lease clauses")
    sr.italic = True
    sr.font.size = Pt(10)
    sr.font.color.rgb = _C_META
    sp.paragraph_format.space_after = Pt(8)

    _hr(doc, "1E40AF")

    def meta(label, value):
        p = doc.add_paragraph()
        l = p.add_run(f"{label}  ")
        l.bold = True
        l.font.size = Pt(9)
        l.font.color.rgb = _C_META
        v = p.add_run(_sanitize(str(value)))
        v.font.size = Pt(9)
        v.font.color.rgb = _C_BODY
        p.paragraph_format.space_after = Pt(2)

    meta("Document:", document_name)
    meta("Generated:", now)
    meta("Total Findings:", str(len(clauses)))
    meta("Analysis Engine:", "Google Gemini")

    _hr(doc, "BFDBFE")
    doc.add_paragraph()

    # ── Executive summary ─────────────────────────────────────────────────────
    h = doc.add_paragraph()
    hh = h.add_run("EXECUTIVE SUMMARY")
    hh.bold = True
    hh.font.size = Pt(13)
    hh.font.color.rgb = _C_TITLE
    h.paragraph_format.space_after = Pt(6)

    high_count   = sum(1 for c in clauses if c.get("severity") == "HIGH")
    medium_count = sum(1 for c in clauses if c.get("severity") == "MEDIUM")
    low_count    = sum(1 for c in clauses if c.get("severity") == "LOW")

    cat_counts = {cat: sum(1 for c in clauses if c.get("category") == cat)
                  for cat in _CATEGORY_ORDER}

    sev_p = doc.add_paragraph()
    sev_p.paragraph_format.space_after = Pt(2)
    sev_p.paragraph_format.left_indent = Pt(12)
    for sev_label, count, color in [
        ("HIGH", high_count, _C_HIGH),
        ("MEDIUM", medium_count, _C_MEDIUM),
        ("LOW", low_count, _C_LOW),
    ]:
        run = sev_p.add_run(f"  {_SEVERITY_ICONS[sev_label]} {count} {sev_label}    ")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = color

    doc.add_paragraph()

    for cat in _CATEGORY_ORDER:
        if cat_counts[cat] == 0:
            continue
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(1)
        cp.paragraph_format.left_indent = Pt(16)
        cr = cp.add_run(f"  •  {_CATEGORY_TITLES[cat]}:  {cat_counts[cat]} finding(s)")
        cr.font.size = Pt(9.5)
        cr.font.color.rgb = _C_DESC

    doc.add_paragraph()

    disc_p = doc.add_paragraph()
    disc_p.paragraph_format.left_indent = Pt(12)
    disc_r = disc_p.add_run(
        "⚠ Note: This analysis is AI-generated and should be reviewed by qualified legal counsel. "
        "Page numbers are approximate and may vary depending on PDF formatting."
    )
    disc_r.italic = True
    disc_r.font.size = Pt(8.5)
    disc_r.font.color.rgb = _C_LABEL

    doc.add_paragraph()
    _hr(doc, "BFDBFE")
    doc.add_paragraph()

    # ── Detailed sections ─────────────────────────────────────────────────────
    for cat in _CATEGORY_ORDER:
        cat_clauses = [c for c in clauses if c.get("category") == cat]
        if not cat_clauses:
            continue

        sh = doc.add_paragraph()
        sh_run = sh.add_run(_CATEGORY_TITLES[cat])
        sh_run.bold = True
        sh_run.font.size = Pt(14)
        sh_run.font.color.rgb = _C_SECTION
        sh.paragraph_format.space_after = Pt(2)

        sd = doc.add_paragraph()
        sd_run = sd.add_run(_CATEGORY_DESCRIPTIONS[cat])
        sd_run.italic = True
        sd_run.font.size = Pt(9)
        sd_run.font.color.rgb = _C_LABEL
        sd.paragraph_format.space_after = Pt(8)

        for idx, clause in enumerate(cat_clauses, 1):
            severity       = clause.get("severity", "MEDIUM")
            title          = _sanitize(clause.get("clause_title", "Untitled"))
            clause_text    = _sanitize(clause.get("clause_text", ""))
            risk_summary   = _sanitize(clause.get("risk_summary", ""))
            page_num       = clause.get("page_number")
            recommendation = _sanitize(clause.get("recommendation", ""))
            sev_color      = _SEVERITY_COLORS.get(severity, _C_MEDIUM)
            sev_icon       = _SEVERITY_ICONS.get(severity, "🟡")

            tp = doc.add_paragraph()
            tp.paragraph_format.space_before = Pt(6)
            tp.paragraph_format.space_after = Pt(2)
            tp.paragraph_format.left_indent = Pt(12)
            badge = tp.add_run(f"{sev_icon} [{severity}]  ")
            badge.bold = True
            badge.font.size = Pt(9.5)
            badge.font.color.rgb = sev_color
            tr = tp.add_run(f"{idx}. {title}")
            tr.bold = True
            tr.font.size = Pt(10.5)
            tr.font.color.rgb = _C_BODY
            if page_num is not None:
                pr = tp.add_run(f"  (Page {page_num})")
                pr.font.size = Pt(9)
                pr.font.color.rgb = _C_LABEL

            if clause_text:
                ct_label = doc.add_paragraph()
                ct_label.paragraph_format.space_before = Pt(2)
                ct_label.paragraph_format.space_after = Pt(1)
                ct_label.paragraph_format.left_indent = Pt(24)
                cl = ct_label.add_run("Clause Text:")
                cl.bold = True
                cl.font.size = Pt(8.5)
                cl.font.color.rgb = _C_LABEL
                ct_p = doc.add_paragraph()
                ct_p.paragraph_format.space_before = Pt(1)
                ct_p.paragraph_format.space_after = Pt(3)
                ct_p.paragraph_format.left_indent = Pt(28)
                ct_p.paragraph_format.right_indent = Pt(20)
                ct_r = ct_p.add_run(f'"{clause_text}"')
                ct_r.italic = True
                ct_r.font.size = Pt(8.5)
                ct_r.font.color.rgb = _C_CLAUSE_BG

            if risk_summary:
                rs_label = doc.add_paragraph()
                rs_label.paragraph_format.space_before = Pt(2)
                rs_label.paragraph_format.space_after = Pt(1)
                rs_label.paragraph_format.left_indent = Pt(24)
                rl = rs_label.add_run("⚠ Why This Matters:")
                rl.bold = True
                rl.font.size = Pt(8.5)
                rl.font.color.rgb = sev_color
                rs_p = doc.add_paragraph()
                rs_p.paragraph_format.space_before = Pt(1)
                rs_p.paragraph_format.space_after = Pt(3)
                rs_p.paragraph_format.left_indent = Pt(28)
                rs_r = rs_p.add_run(risk_summary)
                rs_r.font.size = Pt(9)
                rs_r.font.color.rgb = _C_DESC

            if recommendation:
                rec_label = doc.add_paragraph()
                rec_label.paragraph_format.space_before = Pt(1)
                rec_label.paragraph_format.space_after = Pt(1)
                rec_label.paragraph_format.left_indent = Pt(24)
                recl = rec_label.add_run("💡 Recommendation:")
                recl.bold = True
                recl.font.size = Pt(8.5)
                recl.font.color.rgb = _C_REC
                rec_p = doc.add_paragraph()
                rec_p.paragraph_format.space_before = Pt(1)
                rec_p.paragraph_format.space_after = Pt(4)
                rec_p.paragraph_format.left_indent = Pt(28)
                rec_r = rec_p.add_run(recommendation)
                rec_r.font.size = Pt(9)
                rec_r.font.color.rgb = _C_REC

            _hr(doc, "E5E7EB")

        doc.add_paragraph()

    # ── Footer ────────────────────────────────────────────────────────────────
    _hr(doc, "1E40AF")
    fp = doc.add_paragraph()
    fr = fp.add_run(
        f"End of report  ·  {len(clauses)} findings  ·  "
        f"{high_count} High  ·  {medium_count} Medium  ·  {low_count} Low  ·  {now}"
    )
    fr.font.size = Pt(8)
    fr.font.color.rgb = _C_LABEL

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# MCP TOOLS
# ─────────────────────────────────────────────────────────────────────────────

@mcp.tool()
def build_rent_roll(pdf_bytes: bytes) -> bytes:
    """
    Extract contract data from a lease deed PDF and build a 41-column rent roll Excel file.
    Returns xlsx bytes, with date formatting, percentage columns, and frozen panes.
    Use this when you only need the rent roll without clause analysis.
    """
    if not pdf_bytes:
        raise ValueError("pdf_bytes cannot be empty.")
    contracts = _extract_contracts(pdf_bytes)
    if not contracts:
        raise ValueError("No contracts found in the document.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rent Roll"

    hfill  = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    hfont  = Font(bold=True, color="FFFFFF", size=10, name="Arial")
    dfont  = Font(size=10, name="Arial")
    thin   = Side(style="thin")
    bdr    = Border(left=thin, right=thin, top=thin, bottom=thin)
    calign = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lalign = Alignment(horizontal="left",   vertical="center", wrap_text=False)

    for c, h in enumerate(_HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill      = hfill
        cell.font      = hfont
        cell.alignment = calign
        cell.border    = bdr

    for row_idx, contract in enumerate(contracts, 2):
        contract = _enrich_contract(contract)
        for c in range(1, 42):
            key = _COL_MAP.get(c)
            val = contract.get(key) if key else None
            if c in _DATE_COLS:
                val = _parse_date(val)
            cell = ws.cell(row=row_idx, column=c, value=val)
            cell.font      = dfont
            cell.border    = bdr
            cell.alignment = lalign
            if c in _DATE_COLS and isinstance(val, datetime):
                cell.number_format = "dd-mmm-yy"
            elif c in _PCT_COLS:
                cell.number_format = "0.00%"
            elif c in _PSF_COLS:
                cell.number_format = "0.00"
            elif c in _CURRENCY_COLS:
                cell.number_format = "##,##0"
            elif c in _YRS_COLS:
                cell.number_format = "0.00"
            elif c in _INT_COLS:
                cell.number_format = "0"

    for c, w in enumerate(_COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "D2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


@mcp.tool()
def build_clause_report(pdf_bytes: bytes) -> bytes:
    """
    Analyse critical clauses in a lease deed PDF and build a formatted Word document.
    Clauses are grouped and colour-coded by risk category.
    Returns docx bytes.
    Use this when you only need the clause report without the rent roll.
    """
    if not pdf_bytes:
        raise ValueError("pdf_bytes cannot be empty.")
    clauses = _analyze_clauses(pdf_bytes)
    if not clauses:
        raise ValueError("No critical clauses identified in the document.")
    return _build_docx(clauses)


@mcp.tool()
def process_lease_deed(pdf_bytes: bytes) -> dict:
    """
    Full end-to-end lease deed review. Produces both a rent roll Excel and a
    critical clause report Word document in one call.
    Returns: { "rent_roll_excel": bytes, "clause_report_docx": bytes }
    Use build_rent_roll or build_clause_report individually for partial workflows.
    """
    if not pdf_bytes:
        raise ValueError("pdf_bytes cannot be empty.")
    excel = build_rent_roll(pdf_bytes)
    docx  = build_clause_report(pdf_bytes)
    return {"rent_roll_excel": excel, "clause_report_docx": docx}

"""
Tests for the tax module.

Runs against the real SBI fixture file and a temporary SQLite DB built to
match the base agent's expected schema. No FastAPI / network involved.
"""
import json
import os
import sqlite3
import sys
import tempfile
from pathlib import Path

# Make `tax` importable
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core.schema import migrate
from core.parsers import sbi_excel as sbi_parser
from core.ingest import excel as ingest_excel
from core.classify import manual as classifier
from core.output import excel_export as exporter

FIXTURE = ROOT / "tests" / "fixtures" / "SBI_Transactions_2.xls"


# Replicates the base agent's payments schema closely enough for testing.
BASE_SCHEMA = """
CREATE TABLE IF NOT EXISTS payments (
  id                INTEGER PRIMARY KEY AUTOINCREMENT,
  message_rowid     INTEGER UNIQUE,
  occurred_at       TEXT    NOT NULL,
  amount            REAL    NOT NULL,
  direction         TEXT    NOT NULL,
  bank              TEXT,
  counterparty_vpa  TEXT,
  counterparty_name TEXT,
  ref_no            TEXT,
  sender_id         TEXT,
  raw_text          TEXT,
  confidence        REAL,
  calendar_events   TEXT,
  category          TEXT
);
"""


def _fresh_db() -> Path:
    """Build a temp DB with the base schema, then run our migration."""
    path = Path(tempfile.mkstemp(suffix=".db")[1])
    conn = sqlite3.connect(path)
    conn.executescript(BASE_SCHEMA)
    conn.commit()
    conn.close()
    migrate.migrate(path)
    return path


def _seed_sms_payment(db: Path, ref_no: str, **overrides) -> int:
    conn = sqlite3.connect(db)
    defaults = dict(
        occurred_at="2025-09-01",
        amount=100.0, direction="debit", bank="SBI",
        counterparty_vpa="seed@test", counterparty_name="Seed",
        ref_no=ref_no, source="sms", txn_kind="upi",
    )
    defaults.update(overrides)
    cols = ",".join(defaults.keys())
    qs = ",".join("?" for _ in defaults)
    cur = conn.execute(
        f"INSERT INTO payments ({cols}) VALUES ({qs})", list(defaults.values()))
    conn.commit()
    pid = cur.lastrowid
    conn.close()
    return pid


def _add_bucket(db: Path, name: str):
    conn = sqlite3.connect(db)
    conn.execute(
        "INSERT INTO buckets (name, created_at) VALUES (?, datetime('now'))",
        (name,))
    conn.commit()
    conn.close()


# -------------- Parser tests --------------

def test_sbi_parser_real_file():
    rows = sbi_parser.parse_sbi_file(FIXTURE)
    assert len(rows) >= 700, f"Expected ≥700 rows, got {len(rows)}"
    with_ref = sum(1 for r in rows if r.ref_no)
    assert with_ref >= 680, f"Expected ≥680 with ref_no, got {with_ref}"

    valid_kinds = {"upi", "neft", "imps", "rtgs", "interest", "charges",
                   "insurance", "pos", "atm", "reversal", "other"}
    for r in rows:
        assert r.txn_kind in valid_kinds, f"Bad kind: {r.txn_kind}"
        assert r.direction in ("debit", "credit")
        assert r.amount > 0  # SBI never gives 0-amount rows
        # Date is ISO
        assert len(r.occurred_at) == 10 and r.occurred_at[4] == "-"

    # UPI rows should have ref_no
    upi_with_ref = [r for r in rows if r.txn_kind == "upi" and r.ref_no]
    assert len(upi_with_ref) >= 680, "Most UPI rows must produce ref_no"
    print(f"  ✓ Parsed {len(rows)} rows, {with_ref} with ref_no")


def test_parser_extracts_vpa_and_name():
    rows = sbi_parser.parse_sbi_file(FIXTURE)
    upi = [r for r in rows if r.txn_kind == "upi"]
    with_meta = [r for r in upi if r.counterparty_vpa and r.counterparty_name]
    # Most UPI rows should have full meta
    assert len(with_meta) >= len(upi) * 0.7, \
        f"Only {len(with_meta)}/{len(upi)} UPI rows have full meta"
    print(f"  ✓ {len(with_meta)}/{len(upi)} UPI rows have VPA + name")


# -------------- Dedup tests --------------

def test_dedup_promotes_source():
    """Existing SMS row + same ref_no in Excel → source flips to 'both'."""
    db = _fresh_db()
    try:
        # Find a real UPI ref from the file
        rows = sbi_parser.parse_sbi_file(FIXTURE)
        sample = next(r for r in rows if r.txn_kind == "upi" and r.ref_no)

        sms_id = _seed_sms_payment(db, sample.ref_no)

        stats = ingest_excel.ingest(FIXTURE, db_path=db)

        conn = sqlite3.connect(db); conn.row_factory = sqlite3.Row
        row = conn.execute("SELECT source FROM payments WHERE id = ?",
                           (sms_id,)).fetchone()
        conn.close()
        assert row["source"] == "both", f"source={row['source']}"
        assert stats["matched_existing"] >= 1
        print(f"  ✓ Promoted source to 'both'; total matched: {stats['matched_existing']}")
    finally:
        os.unlink(db)


def test_dedup_inserts_neft_as_excel_only():
    """NEFT rows have no SMS counterpart → all inserted with source='excel'."""
    db = _fresh_db()
    try:
        ingest_excel.ingest(FIXTURE, db_path=db)
        conn = sqlite3.connect(db); conn.row_factory = sqlite3.Row
        neft = conn.execute(
            "SELECT source, txn_kind FROM payments WHERE txn_kind = 'neft'"
        ).fetchall()
        conn.close()
        assert len(neft) >= 16, f"Expected ≥16 NEFT, got {len(neft)}"
        for r in neft:
            assert r["source"] == "excel"
        print(f"  ✓ {len(neft)} NEFT rows inserted as source='excel'")
    finally:
        os.unlink(db)


def test_reingest_is_idempotent():
    """Running ingest twice on the same file should not double-insert."""
    db = _fresh_db()
    try:
        s1 = ingest_excel.ingest(FIXTURE, db_path=db)
        s2 = ingest_excel.ingest(FIXTURE, db_path=db)

        # Second run should match all rows that produced ref_no, insert 0
        # for those, but may insert ref_no-less rows again (acceptable
        # tradeoff documented in plan).
        assert s2["matched_existing"] >= s1["inserted"] - 10, (
            f"Re-ingest matched {s2['matched_existing']}, "
            f"first inserted {s1['inserted']}"
        )
        print(f"  ✓ Re-ingest matched {s2['matched_existing']} of {s1['inserted']}")
    finally:
        os.unlink(db)


# -------------- Classifier tests --------------

def test_classifier_vpa_exact_match():
    db = _fresh_db()
    try:
        conn = sqlite3.connect(db)
        conn.execute(
            "INSERT INTO category_rules (key_kind, key_value, category, hits) "
            "VALUES ('vpa', 'zepto.payu', 'food', 5)")
        conn.commit()

        clf = classifier.RuleBasedClassifier(conn)
        s = clf.suggest({"counterparty_vpa": "zepto.payu",
                         "counterparty_name": "Zepto Marketplace"})
        assert s.bucket == "food"
        assert s.confidence == 1.0
        conn.close()
        print(f"  ✓ VPA exact match → conf 1.0, reason: {s.reason}")
    finally:
        os.unlink(db)


def test_classifier_substring_match():
    db = _fresh_db()
    try:
        conn = sqlite3.connect(db)
        conn.execute(
            "INSERT INTO category_rules (key_kind, key_value, category, hits) "
            "VALUES ('name', 'Zepto', 'food', 3)")
        conn.commit()

        clf = classifier.RuleBasedClassifier(conn)
        s = clf.suggest({"counterparty_vpa": "different@vpa",
                         "counterparty_name": "Zepto Marketplace Ltd"})
        assert s.bucket == "food"
        assert s.confidence == 0.6
        conn.close()
        print(f"  ✓ Substring match → conf 0.6")
    finally:
        os.unlink(db)


def test_classifier_no_match():
    db = _fresh_db()
    try:
        conn = sqlite3.connect(db)
        clf = classifier.RuleBasedClassifier(conn)
        s = clf.suggest({"counterparty_vpa": "unknown@bank",
                         "counterparty_name": "Random Merchant"})
        assert s.bucket is None
        assert s.confidence == 0.0
        conn.close()
        print(f"  ✓ No match → confidence 0.0")
    finally:
        os.unlink(db)


def test_classifier_learn_and_recall():
    db = _fresh_db()
    try:
        conn = sqlite3.connect(db)
        clf = classifier.RuleBasedClassifier(conn)
        payment = {"counterparty_vpa": "uber.driver@paytm",
                   "counterparty_name": "Uber India"}
        clf.learn(payment, "transport")
        clf.learn(payment, "transport")  # second call should bump hits

        rule = conn.execute(
            "SELECT category, hits FROM category_rules "
            "WHERE key_kind='vpa' AND key_value='uber.driver@paytm'"
        ).fetchone()
        assert rule[0] == "transport"
        assert rule[1] == 2

        s = clf.suggest(payment)
        assert s.bucket == "transport"
        assert s.confidence == 1.0
        assert "2 prior hits" in s.reason
        conn.close()
        print(f"  ✓ Learned + recalled with hits=2")
    finally:
        os.unlink(db)


def test_local_llm_falls_back_when_unreachable():
    db = _fresh_db()
    try:
        conn = sqlite3.connect(db)
        # No category_rules; LLM endpoint unreachable
        clf = classifier.LocalLLMClassifier(
            conn, endpoint="http://127.0.0.1:1",  # closed port
            model="dummy", buckets=["food", "transport"], timeout=0.5)
        s = clf.suggest({"counterparty_vpa": "x@y",
                         "counterparty_name": "Z"})
        # Should fall back to rule-based (which returns None here)
        assert s.bucket is None  # no rule, no reachable LLM
        conn.close()
        print(f"  ✓ Unreachable LLM → graceful rule-based fallback")
    finally:
        os.unlink(db)


def test_active_classifier_picks_correct_impl():
    db = _fresh_db()
    try:
        # mode='manual' (default) → RuleBasedClassifier
        conn = sqlite3.connect(db)
        clf = classifier.get_active_classifier(conn)
        assert isinstance(clf, classifier.RuleBasedClassifier)

        # mode='ai' but no endpoint → still RuleBasedClassifier
        conn.execute("UPDATE tax_settings SET value='ai' WHERE key='mode'")
        conn.commit()
        clf = classifier.get_active_classifier(conn)
        assert isinstance(clf, classifier.RuleBasedClassifier)

        # mode='ai' + endpoint set + buckets exist → LocalLLMClassifier
        conn.execute(
            "UPDATE tax_settings SET value='http://127.0.0.1:11434' "
            "WHERE key='llm_endpoint'")
        conn.execute(
            "INSERT INTO buckets (name, created_at) VALUES "
            "('food', datetime('now'))")
        conn.commit()
        clf = classifier.get_active_classifier(conn)
        assert isinstance(clf, classifier.LocalLLMClassifier)
        conn.close()
        print(f"  ✓ Factory picks correct implementation per mode")
    finally:
        os.unlink(db)


# -------------- Retroactive classification (simulating routes logic) --------

def test_retro_application():
    """Simulate the route's apply_retro=true logic on multiple unclassified rows."""
    db = _fresh_db()
    try:
        # Ingest the real file → many Zepto rows (or whatever common merchant)
        ingest_excel.ingest(FIXTURE, db_path=db)
        _add_bucket(db, "food")

        conn = sqlite3.connect(db); conn.row_factory = sqlite3.Row

        # Find a counterparty_name with multiple unclassified rows
        candidate = conn.execute("""
            SELECT counterparty_name, COUNT(*) AS n FROM payments
            WHERE counterparty_name IS NOT NULL
              AND (category IS NULL OR category = '')
            GROUP BY counterparty_name HAVING n >= 3
            ORDER BY n DESC LIMIT 1
        """).fetchone()
        assert candidate is not None, "No merchant with ≥3 unclassified rows"
        merchant = candidate["counterparty_name"]
        n_total = candidate["n"]

        # Classify ONE of those rows + apply retroactively
        first = conn.execute(
            "SELECT * FROM payments WHERE counterparty_name = ? "
            "  AND (category IS NULL OR category = '') LIMIT 1",
            (merchant,)).fetchone()

        conn.execute(
            "UPDATE payments SET category='food', category_source='user' "
            "WHERE id = ?", (first["id"],))
        retro = conn.execute("""
            UPDATE payments
            SET category = 'food', category_source = 'auto'
            WHERE (category IS NULL OR category = '')
              AND id != ?
              AND counterparty_name = ?
        """, (first["id"], merchant)).rowcount
        conn.commit()

        all_rows = conn.execute(
            "SELECT category FROM payments WHERE counterparty_name = ?",
            (merchant,)).fetchall()
        all_food = all([r["category"] == "food" for r in all_rows])
        conn.close()

        assert retro == n_total - 1, f"Expected {n_total-1} retro, got {retro}"
        assert all_food, "Not all merchant rows classified after retro"
        print(f"  ✓ Retro applied to {retro} other '{merchant}' rows")
    finally:
        os.unlink(db)


def test_bucket_rename_cascades():
    db = _fresh_db()
    try:
        _add_bucket(db, "food")
        conn = sqlite3.connect(db)
        # Insert a payment + a category rule using the bucket name
        conn.execute("""
            INSERT INTO payments (occurred_at, amount, direction, category)
            VALUES ('2025-10-01', 100.0, 'debit', 'food')""")
        conn.execute("""
            INSERT INTO category_rules (key_kind, key_value, category, hits)
            VALUES ('vpa', 'zepto.payu', 'food', 1)""")
        conn.commit()

        # Simulate the route's rename cascade
        bid = conn.execute("SELECT id FROM buckets WHERE name='food'").fetchone()[0]
        old = "food"; new = "groceries"
        conn.execute("UPDATE buckets SET name=? WHERE id=?", (new, bid))
        conn.execute("UPDATE payments SET category=? WHERE category=?", (new, old))
        conn.execute("UPDATE category_rules SET category=? WHERE category=?", (new, old))
        conn.commit()

        p_cat = conn.execute("SELECT category FROM payments LIMIT 1").fetchone()[0]
        r_cat = conn.execute("SELECT category FROM category_rules LIMIT 1").fetchone()[0]
        conn.close()
        assert p_cat == "groceries"
        assert r_cat == "groceries"
        print(f"  ✓ Rename 'food' → 'groceries' cascaded to payments + rules")
    finally:
        os.unlink(db)


# -------------- Export tests --------------

def test_export_xlsx_has_formulas():
    db = _fresh_db()
    try:
        _add_bucket(db, "food")
        ingest_excel.ingest(FIXTURE, db_path=db)
        # Classify a handful so Summary has content
        conn = sqlite3.connect(db)
        conn.execute("""
            UPDATE payments SET category='food', category_source='user'
            WHERE id IN (SELECT id FROM payments
                         WHERE category IS NULL OR category=''
                         ORDER BY id LIMIT 5)""")
        conn.commit(); conn.close()

        out_path = exporter.export_xlsx(db_path=db)
        assert Path(out_path).exists()

        # Re-open with formulas (data_only=False) and verify the SUMIFS
        from openpyxl import load_workbook
        wb = load_workbook(out_path, data_only=False)
        assert "Transactions" in wb.sheetnames
        assert "Summary" in wb.sheetnames

        s = wb["Summary"]
        # Find a non-Total bucket row and check B has SUMIFS
        formula_found = False
        for r in range(2, s.max_row + 1):
            v = s.cell(row=r, column=2).value
            if isinstance(v, str) and v.startswith("=SUMIFS"):
                formula_found = True; break
        assert formula_found, "No SUMIFS formula in Summary sheet"

        # Transactions sheet has frozen pane
        t = wb["Transactions"]
        assert t.freeze_panes == "A2"
        Path(out_path).unlink()
        print(f"  ✓ Export has SUMIFS formulas + frozen header")
    finally:
        os.unlink(db)


# -------------- Run --------------

def run_all():
    tests = [v for k, v in globals().items()
             if k.startswith("test_") and callable(v)]
    print(f"Running {len(tests)} tests...\n")
    failed = []
    for t in tests:
        try:
            print(f"• {t.__name__}")
            t()
        except AssertionError as e:
            failed.append((t.__name__, str(e)))
            print(f"  ✗ FAIL: {e}")
        except Exception as e:
            failed.append((t.__name__, f"{type(e).__name__}: {e}"))
            print(f"  ✗ ERROR: {type(e).__name__}: {e}")

    print(f"\n{'─'*60}")
    if failed:
        print(f"FAILED: {len(failed)}/{len(tests)}")
        for name, err in failed:
            print(f"  {name}: {err}")
        sys.exit(1)
    else:
        print(f"PASSED: {len(tests)}/{len(tests)}")


if __name__ == "__main__":
    run_all()

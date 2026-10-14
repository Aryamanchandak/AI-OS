"""
Export classified transactions to a single annotated .xlsx file.

Two sheets:
  - Transactions: every txn with bucket assignment, frozen header, autofilter
  - Summary: bucket totals using SUMIFS formulas (live, not pre-computed)

The user hands this to their CA. The CA can resort, refilter, and the
Summary recalculates because we use formulas, not hardcoded numbers.
"""
from __future__ import annotations

import sqlite3
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from core import config


def export_xlsx(db_path: Path | None = None, out_path: Path | None = None) -> str:
    db_path = db_path or config.DB_PATH
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    rows = conn.execute("""
        SELECT
          occurred_at,
          COALESCE(txn_kind, 'upi')   AS txn_kind,
          direction,
          amount,
          COALESCE(counterparty_name, '') AS counterparty,
          COALESCE(counterparty_vpa, '')  AS vpa,
          COALESCE(ref_no, '')            AS ref_no,
          COALESCE(category, '(unclassified)') AS bucket,
          COALESCE(source, 'sms') AS source
        FROM payments
        ORDER BY occurred_at ASC, id ASC
    """).fetchall()
    conn.close()

    wb = Workbook()

    ws = wb.active
    ws.title = "Transactions"

    headers = ["Date", "Kind", "Direction", "Amount", "Counterparty",
               "VPA", "Ref No", "Bucket", "Source"]
    ws.append(headers)
    header_font = Font(bold=True, name="Arial")
    header_fill = PatternFill("solid", start_color="EEEEEE")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")

    for r in rows:
        ws.append([
            r["occurred_at"],
            r["txn_kind"],
            r["direction"],
            float(r["amount"]) if r["amount"] is not None else None,
            r["counterparty"],
            r["vpa"],
            r["ref_no"],
            r["bucket"],
            r["source"],
        ])

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 32
    ws.column_dimensions["F"].width = 24
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 18
    ws.column_dimensions["I"].width = 10

    for cell in ws["D"][1:]:
        cell.number_format = '"₹"#,##0.00;[Red]("₹"#,##0.00);-'
        cell.font = Font(name="Arial")

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("Summary")
    s.append(["Bucket", "Debits", "Credits", "Net", "Count"])
    for cell in s[1]:
        cell.font = header_font
        cell.fill = header_fill

    buckets = sorted({(r["bucket"]) for r in rows},
                     key=lambda b: (b == "(unclassified)", b))
    last_data_row = ws.max_row

    for i, b in enumerate(buckets, start=2):
        s.cell(row=i, column=1, value=b)
        s.cell(row=i, column=2, value=(
            f'=SUMIFS(Transactions!D2:D{last_data_row},'
            f'Transactions!H2:H{last_data_row},A{i},'
            f'Transactions!C2:C{last_data_row},"debit")'
        ))
        s.cell(row=i, column=3, value=(
            f'=SUMIFS(Transactions!D2:D{last_data_row},'
            f'Transactions!H2:H{last_data_row},A{i},'
            f'Transactions!C2:C{last_data_row},"credit")'
        ))
        s.cell(row=i, column=4, value=f"=C{i}-B{i}")
        s.cell(row=i, column=5, value=(
            f'=COUNTIF(Transactions!H2:H{last_data_row},A{i})'
        ))

    if buckets:
        total_row = len(buckets) + 2
        s.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
        s.cell(row=total_row, column=2, value=f"=SUM(B2:B{total_row-1})").font = Font(bold=True)
        s.cell(row=total_row, column=3, value=f"=SUM(C2:C{total_row-1})").font = Font(bold=True)
        s.cell(row=total_row, column=4, value=f"=SUM(D2:D{total_row-1})").font = Font(bold=True)
        s.cell(row=total_row, column=5, value=f"=SUM(E2:E{total_row-1})").font = Font(bold=True)

    for col in ("B", "C", "D"):
        for cell in s[col][1:]:
            cell.number_format = '"₹"#,##0.00;[Red]("₹"#,##0.00);-'
            if cell.font.bold:
                cell.font = Font(name="Arial", bold=True)
            else:
                cell.font = Font(name="Arial")

    s.column_dimensions["A"].width = 22
    s.column_dimensions["B"].width = 16
    s.column_dimensions["C"].width = 16
    s.column_dimensions["D"].width = 16
    s.column_dimensions["E"].width = 10
    s.freeze_panes = "A2"

    if out_path is None:
        out_path = Path(tempfile.gettempdir()) / (
            f"transactions-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
        )
    wb.save(out_path)
    return str(out_path)


if __name__ == "__main__":
    print(export_xlsx())
